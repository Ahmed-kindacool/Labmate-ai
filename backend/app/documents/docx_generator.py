"""Phase 7: renders a `GeneratedLab` + execution results + screenshots into
a final, downloadable `.docx`, using the university's template from
`TemplateRegistry` (Phase 6).

This is the last stage of the pipeline `AI_AND_GENERATION.md` describes:

    Lab Parser -> AI Service -> Code Executor -> Screenshot Service
    -> Template Registry -> DOCX Generator (this module)

Design note on the {{TASKS}} / {{CODE}} / {{OUTPUT_SCREENSHOT}} split:
the template (Phase 6) has one top-level section per placeholder, matching
TEMPLATES_AND_UI.md's flat placeholder list. Rather than reshaping the
template into a single repeated "task card" block, this generator expands
each placeholder into its own list, one entry per task, keeping the
template's existing three-section shape (Tasks, then Code, then Execution
Output) and just filling each section with every task's content in order.
"""

from __future__ import annotations

import base64
import binascii
import io

from docx import Document
from docx.enum.text import WD_BREAK
from docx.image.exceptions import UnrecognizedImageError
from docx.shared import Inches, Pt

from app.core.errors import AppError
from app.documents.docx_helpers import (
    find_placeholder_paragraph,
    insert_paragraph_after,
    iter_all_paragraphs,
    remove_paragraph,
    replace_placeholder_text,
)
from app.schemas.generate import TaskExecutionResult
from app.schemas.lab import ExecutionResult, ExecutionStatus, GeneratedLab, GeneratedTaskSolution
from app.schemas.student import StudentInfo
from app.templates_registry import TemplateRegistry

_CODE_FONT = "Consolas"
_CODE_SIZE = Pt(10)
_SCREENSHOT_WIDTH = Inches(5.5)

# Statuses for which the AI/executor pipeline produced real, capturable
# output text (stdout on success, stderr as the actual error otherwise).
# Mirrors app/screenshots/service.py's _CAPTURABLE_STATUSES -- kept
# separate rather than imported, since "what's worth a raw-text fallback
# here" and "what's worth a screenshot there" are similar but independent
# decisions.
_TEXT_FALLBACK_STATUSES = {ExecutionStatus.SUCCESS, ExecutionStatus.FAILED, ExecutionStatus.TIMEOUT}


class DocxGenerator:
    """Fills a university's template.docx with one generated lab report."""

    def __init__(self, template_registry: TemplateRegistry | None = None) -> None:
        self._templates = template_registry or TemplateRegistry()

    def generate(
        self,
        student: StudentInfo,
        generated_lab: GeneratedLab,
        execution_results: list[TaskExecutionResult] | None = None,
        screenshots: dict[str, str] | None = None,
    ) -> bytes:
        """Returns the finished report as raw `.docx` bytes.

        Raises `AppError` (code `REPORT_GENERATION_FAILED`) if the
        template can't be found/opened, or if assembling the final
        document fails -- both per `docs/API_CONTRACT.md`'s error table.
        `TemplateRegistry.get_template` already raises `AppError` itself
        for a missing/unrecognized template, so that case passes through
        unchanged rather than being re-wrapped.
        """
        execution_results = execution_results or []
        screenshots = screenshots or {}
        results_by_task = {r.task_id: r.result for r in execution_results}

        template_config = self._templates.get_template(student.university)
        try:
            document = Document(str(template_config.template_path))
        except Exception as exc:  # noqa: BLE001 - any docx/zip corruption, wrapped below
            raise AppError(
                "REPORT_GENERATION_FAILED",
                "Could not open the report template for this university.",
            ) from exc

        self._fill_student_and_lab_fields(document, student, generated_lab)
        self._fill_objectives(document, generated_lab.objectives)
        self._fill_tasks(document, generated_lab.tasks)
        self._fill_code(document, generated_lab.tasks)
        self._fill_output(document, generated_lab.tasks, results_by_task, screenshots)

        buffer = io.BytesIO()
        try:
            document.save(buffer)
        except Exception as exc:  # noqa: BLE001
            raise AppError(
                "REPORT_GENERATION_FAILED",
                "Could not assemble the final report document.",
            ) from exc
        return buffer.getvalue()

    # -- simple, single-value placeholders --------------------------------

    def _fill_student_and_lab_fields(
        self, document: Document, student: StudentInfo, generated_lab: GeneratedLab
    ) -> None:
        values = {
            "{{STUDENT_NAME}}": student.name,
            "{{ROLL_NUMBER}}": student.roll_number,
            "{{CLASS_SECTION}}": student.class_section,
            "{{INSTRUCTOR_NAME}}": student.instructor_name,
            "{{COURSE}}": student.course,
            "{{LAB_TITLE}}": generated_lab.lab_title,
            "{{CONCLUSION}}": generated_lab.conclusion,
        }
        for paragraph in iter_all_paragraphs(document):
            for token, value in values.items():
                replace_placeholder_text(paragraph, token, value)

    # -- {{OBJECTIVES}}: one bullet paragraph per objective ----------------

    def _fill_objectives(self, document: Document, objectives: list[str]) -> None:
        target = find_placeholder_paragraph(document, "{{OBJECTIVES}}")
        if target is None:
            return

        if not objectives:
            replace_placeholder_text(
                target, "{{OBJECTIVES}}", "No objectives were extracted for this lab."
            )
            return

        anchor = target
        for objective in objectives:
            anchor = insert_paragraph_after(anchor)
            anchor.add_run(f"\u2022 {objective}")
        remove_paragraph(target)

    # -- {{TASKS}}: one numbered description per task ----------------------

    def _fill_tasks(self, document: Document, tasks: list[GeneratedTaskSolution]) -> None:
        target = find_placeholder_paragraph(document, "{{TASKS}}")
        if target is None:
            return

        if not tasks:
            replace_placeholder_text(target, "{{TASKS}}", "No tasks were extracted for this lab.")
            return

        anchor = target
        for index, task in enumerate(tasks, start=1):
            anchor = insert_paragraph_after(anchor)
            label_run = anchor.add_run(f"Task {index}: ")
            label_run.bold = True
            anchor.add_run(task.description)
        remove_paragraph(target)

    # -- {{CODE}}: every task's code, each under its own heading -----------

    def _fill_code(self, document: Document, tasks: list[GeneratedTaskSolution]) -> None:
        target = find_placeholder_paragraph(document, "{{CODE}}")
        if target is None:
            return

        if not tasks:
            replace_placeholder_text(target, "{{CODE}}", "No code was generated for this lab.")
            return

        anchor = target
        for index, task in enumerate(tasks, start=1):
            anchor = insert_paragraph_after(anchor)
            heading_run = anchor.add_run(f"Task {index} \u2014 {task.filename}")
            heading_run.bold = True

            anchor = insert_paragraph_after(anchor)
            if task.code.strip():
                self._write_code_block(anchor, task.code)
            else:
                anchor.add_run("(no code required for this task)")
        remove_paragraph(target)

    @staticmethod
    def _write_code_block(paragraph, code: str) -> None:
        run = paragraph.add_run()
        run.font.name = _CODE_FONT
        run.font.size = _CODE_SIZE
        lines = code.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                run.add_break(WD_BREAK.LINE)
            run.add_text(line)

    # -- {{OUTPUT_SCREENSHOT}}: real screenshot, or a text/status fallback -

    def _fill_output(
        self,
        document: Document,
        tasks: list[GeneratedTaskSolution],
        results_by_task: dict[str, ExecutionResult],
        screenshots: dict[str, str],
    ) -> None:
        target = find_placeholder_paragraph(document, "{{OUTPUT_SCREENSHOT}}")
        if target is None:
            return

        if not tasks:
            replace_placeholder_text(
                target, "{{OUTPUT_SCREENSHOT}}", "No tasks were executed for this lab."
            )
            return

        anchor = target
        for index, task in enumerate(tasks, start=1):
            anchor = insert_paragraph_after(anchor)
            heading_run = anchor.add_run(f"Task {index} output:")
            heading_run.bold = True

            anchor = insert_paragraph_after(anchor)
            self._write_task_output(
                anchor, screenshots.get(task.id), results_by_task.get(task.id)
            )
        remove_paragraph(target)

    @staticmethod
    def _write_task_output(
        paragraph, screenshot_b64: str | None, result: ExecutionResult | None
    ) -> None:
        # Per AI_AND_GENERATION.md's anti-fabrication rule, every branch
        # here shows either a real captured artifact (image or raw
        # stdout/stderr text) or an honest status note -- never invented
        # output.
        if screenshot_b64:
            try:
                image_bytes = base64.b64decode(screenshot_b64)
                paragraph.add_run().add_picture(io.BytesIO(image_bytes), width=_SCREENSHOT_WIDTH)
                return
            except (binascii.Error, UnrecognizedImageError, ValueError):
                # Corrupt/undecodable screenshot data. Screenshotting is
                # an enhancement, not core correctness (see
                # GenerationService._capture_screenshots) -- degrade to
                # the raw-text fallback below rather than failing the
                # whole report over one bad image.
                pass

        if result is not None and result.status in _TEXT_FALLBACK_STATUSES:
            text = result.stdout if result.status == ExecutionStatus.SUCCESS else result.stderr
            run = paragraph.add_run(text.strip() or "(no output)")
            run.font.name = _CODE_FONT
            run.font.size = _CODE_SIZE
            return

        if result is not None and result.status == ExecutionStatus.UNSUPPORTED:
            paragraph.add_run(
                "Execution is not supported for this task's language or environment."
            )
            return

        paragraph.add_run("No code was executed for this task.")
