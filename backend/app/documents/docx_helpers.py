"""Low-level python-docx helpers shared by DocxGenerator.

python-docx has no built-in "find this placeholder and replace it with
N paragraphs" operation -- these are the small, generic primitives that
make DocxGenerator's placeholder-by-placeholder methods readable. None of
these know about GeneratedLab, ExecutionResult, or anything else
lab-specific.
"""

from __future__ import annotations

from typing import Iterator

from docx.document import Document as DocumentObject
from docx.text.paragraph import Paragraph


def iter_all_paragraphs(document: DocumentObject) -> Iterator[Paragraph]:
    """Every paragraph in the document body and in table cells.

    Our templates only ever put placeholders in the body or in the
    student-info table, never inside headers/footers (the logo is a
    real image, not a placeholder) -- so this deliberately doesn't walk
    section headers/footers.
    """
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def find_placeholder_paragraph(document: DocumentObject, token: str) -> Paragraph | None:
    """The first paragraph whose combined run text contains `token`, or
    None if no paragraph in the document contains it.

    A missing placeholder is not necessarily a bug -- see each
    DocxGenerator method's handling of a None result -- since a future
    template edit could legitimately drop a section.
    """
    for paragraph in iter_all_paragraphs(document):
        if token in "".join(run.text for run in paragraph.runs):
            return paragraph
    return None


def replace_placeholder_text(paragraph: Paragraph, token: str, value: str) -> bool:
    """Replaces `token` with `value` inside `paragraph`'s combined text,
    in place. Returns True if the token was found and replaced.

    python-docx frequently splits a paragraph's visible text across
    several runs even when the underlying content looks contiguous, so a
    naive `run.text.replace(...)` per-run can silently miss a token split
    across a run boundary. This merges every run's text, does the
    replacement once, and puts the result back in the first run --
    losing any formatting differences between runs after the first,
    which is fine here since every placeholder in our templates is
    authored as a single run to begin with (see docs/TEMPLATE_REGISTRY.md).
    """
    full_text = "".join(run.text for run in paragraph.runs)
    if token not in full_text:
        return False

    new_text = full_text.replace(token, value)
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)
    return True


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    """Inserts and returns a new, empty paragraph immediately after
    `paragraph`, as a sibling in the same body/cell.

    Standard python-docx idiom -- the library only exposes
    `add_paragraph` (always at the end of the body) and
    `insert_paragraph_before`, neither of which lets us build a run of
    new paragraphs at a specific mid-document position while walking
    forward. Operates directly on the underlying XML element.
    """
    from docx.oxml import OxmlElement

    new_p_element = OxmlElement("w:p")
    paragraph._p.addnext(new_p_element)
    return Paragraph(new_p_element, paragraph._parent)


def remove_paragraph(paragraph: Paragraph) -> None:
    """Deletes `paragraph` from the document entirely."""
    element = paragraph._p
    element.getparent().remove(element)
    paragraph._p = None  # type: ignore[assignment]
