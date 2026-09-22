from pathlib import Path

import pytest

from app.core.errors import AppError
from app.domain.university import University
from app.schemas.template import TemplateConfig
from app.templates_registry import TemplateRegistry


class TestGetTemplate:
    """These check against the REAL files checked into
    backend/templates_registry/templates/ -- not mocks -- since the whole
    point of Phase 6 is that those files actually exist and are valid,
    which is exactly what the Phase 2 stub never verified.
    """

    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_returns_real_existing_files_for_each_university(self, university: University) -> None:
        registry = TemplateRegistry()
        config = registry.get_template(university)

        assert isinstance(config, TemplateConfig)
        assert config.university == university
        assert config.template_path.is_file()
        assert config.template_path.name == "template.docx"
        assert config.logo_path.is_file()
        assert config.logo_path.name == "logo.png"

    @pytest.mark.parametrize("raw_value", ["air", "AIR", "Air", "bahria", "NUST"])
    def test_accepts_case_insensitive_string_input(self, raw_value: str) -> None:
        registry = TemplateRegistry()
        config = registry.get_template(raw_value)
        assert config.university.value == raw_value.lower()

    def test_unknown_university_raises_app_error(self) -> None:
        registry = TemplateRegistry()
        with pytest.raises(AppError) as exc_info:
            registry.get_template("mit")

        assert exc_info.value.code == "REPORT_GENERATION_FAILED"
        assert "mit" in exc_info.value.message

    def test_resolved_paths_are_not_relative_to_cwd(self, monkeypatch, tmp_path: Path) -> None:
        """The Phase 2 stub used a hardcoded relative string
        ("backend/templates_registry/...") that only worked if the process
        happened to be launched from the repo root. Changing the CWD here
        must not break resolution.
        """
        monkeypatch.chdir(tmp_path)
        registry = TemplateRegistry()
        config = registry.get_template(University.AIR)
        assert config.template_path.is_file()

    def test_missing_template_file_raises_app_error(self, tmp_path: Path) -> None:
        uni_dir = tmp_path / "air"
        uni_dir.mkdir()
        (uni_dir / "logo.png").write_bytes(b"fake-png")
        # template.docx deliberately not created

        registry = TemplateRegistry(templates_root=tmp_path)
        with pytest.raises(AppError) as exc_info:
            registry.get_template(University.AIR)

        assert exc_info.value.code == "REPORT_GENERATION_FAILED"
        assert "template.docx" in exc_info.value.message

    def test_missing_logo_file_raises_app_error(self, tmp_path: Path) -> None:
        uni_dir = tmp_path / "bahria"
        uni_dir.mkdir()
        (uni_dir / "template.docx").write_bytes(b"fake-docx")
        # logo.png deliberately not created

        registry = TemplateRegistry(templates_root=tmp_path)
        with pytest.raises(AppError) as exc_info:
            registry.get_template(University.BAHRIA)

        assert exc_info.value.code == "REPORT_GENERATION_FAILED"
        assert "logo.png" in exc_info.value.message

    def test_missing_university_directory_entirely_raises_app_error(self, tmp_path: Path) -> None:
        registry = TemplateRegistry(templates_root=tmp_path)
        with pytest.raises(AppError) as exc_info:
            registry.get_template(University.NUST)

        assert exc_info.value.code == "REPORT_GENERATION_FAILED"


class TestRealTemplateContents:
    """Sanity-checks the actual checked-in .docx files, per
    TEMPLATES_AND_UI.md's placeholder list -- catches a template being
    hand-edited into something Phase 7 can no longer find placeholders in.
    """

    PLACEHOLDERS = [
        "{{STUDENT_NAME}}",
        "{{ROLL_NUMBER}}",
        "{{CLASS_SECTION}}",
        "{{INSTRUCTOR_NAME}}",
        "{{COURSE}}",
        "{{LAB_TITLE}}",
        "{{OBJECTIVES}}",
        "{{TASKS}}",
        "{{CODE}}",
        "{{OUTPUT_SCREENSHOT}}",
        "{{CONCLUSION}}",
    ]

    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_template_contains_all_required_placeholders(self, university: University) -> None:
        from docx import Document

        registry = TemplateRegistry()
        config = registry.get_template(university)
        document = Document(str(config.template_path))

        text_chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text_chunks.append(cell.text)
        full_text = "\n".join(text_chunks)

        for placeholder in self.PLACEHOLDERS:
            assert placeholder in full_text, f"{placeholder} missing from {university.value} template"

    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_template_never_contains_university_name_as_text(self, university: University) -> None:
        """Per TEMPLATES_AND_UI.md's "Critical requirement": the selected
        university must never be inserted into the report as ordinary
        text -- only the logo/branding should differ.
        """
        from docx import Document

        registry = TemplateRegistry()
        config = registry.get_template(university)
        document = Document(str(config.template_path))

        text_chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text_chunks.append(cell.text)
        full_text = "\n".join(text_chunks).lower()

        for name in ("air university", "bahria university", "nust", "national university of sciences"):
            assert name not in full_text
