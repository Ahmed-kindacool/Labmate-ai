"""Regression tests for the 'it doesn't actually generate reports' fixes."""

import io
import json
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from app.ai.openai_provider import OpenAIProvider, _parse_generated_lab
from app.core.config import Settings, settings
from app.core.errors import AppError
from app.documents import DocxGenerator
from app.domain.university import University
from app.schemas.lab import GeneratedLab, GeneratedTaskSolution, ParsedLab
from app.schemas.student import StudentInfo

_LAB = {
    "lab_title": "Lab 1",
    "objectives": ["Learn loops"],
    "tasks": [
        {
            "id": "task-1",
            "description": "Print numbers.",
            "language": "python",
            "filename": "a.py",
            "code": "print(1)",
            "explanation": "Prints one.",
        }
    ],
    "conclusion": "Loops are useful.",
}


def _client_returning(*contents):
    """Each call returns the next content (or raises it if it's an exception)."""
    calls = []

    async def create(**kwargs):
        calls.append(kwargs)
        item = contents[min(len(calls), len(contents)) - 1]
        if isinstance(item, Exception):
            raise item
        message = MagicMock()
        message.content = item
        choice = MagicMock()
        choice.message = message
        response = MagicMock()
        response.choices = [choice]
        return response

    client = MagicMock()
    client.chat.completions.create = create
    client.calls = calls
    return client


class _HttpError(Exception):
    def __init__(self, status_code):
        super().__init__(f"http {status_code}")
        self.status_code = status_code


@pytest.fixture
def parsed() -> ParsedLab:
    return ParsedLab(title="t", raw_text="Task 1: print", tasks=[])


class TestAiParsing:
    def test_code_fenced_json_is_accepted(self) -> None:
        lab = _parse_generated_lab("```json\n" + json.dumps(_LAB) + "\n```")
        assert lab.lab_title == "Lab 1"

    def test_prose_around_json_is_ignored(self) -> None:
        lab = _parse_generated_lab("Sure! Here you go:\n" + json.dumps(_LAB) + "\nHope that helps.")
        assert len(lab.tasks) == 1

    def test_missing_optional_fields_are_defaulted_not_fatal(self) -> None:
        sparse = {"lab_title": "L", "tasks": [{"description": "d", "language": "Python", "code": "print(1)"}]}
        lab = _parse_generated_lab(json.dumps(sparse))
        task = lab.tasks[0]
        assert task.id == "task-1"
        assert task.filename == "task1.py"
        assert task.language == "python"
        assert lab.objectives == [] and lab.conclusion == ""

    def test_no_code_task_keeps_empty_filename(self) -> None:
        lab = _parse_generated_lab(json.dumps({"lab_title": "L", "tasks": [{"description": "theory"}]}))
        assert lab.tasks[0].code == "" and lab.tasks[0].filename == ""

    @pytest.mark.parametrize("raw", ["", "no json here", '{"tasks": "nope"}', '{"unexpected": 1}'])
    def test_garbage_still_fails_cleanly(self, raw: str) -> None:
        with pytest.raises(AppError) as exc:
            _parse_generated_lab(raw)
        assert exc.value.code == "AI_GENERATION_FAILED"


class TestProviderBehaviour:
    @pytest.mark.asyncio
    async def test_retries_without_json_mode_when_provider_rejects_it(self, parsed) -> None:
        client = _client_returning(_HttpError(400), json.dumps(_LAB))
        lab = await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert lab.lab_title == "Lab 1"
        assert "response_format" in client.calls[0]
        assert "response_format" not in client.calls[1]

    @pytest.mark.asyncio
    @pytest.mark.parametrize(
        ("status", "fragment"),
        [(401, "API key"), (404, "model"), (429, "busy or out of quota")],
    )
    async def test_failures_give_actionable_messages(self, parsed, status, fragment) -> None:
        client = _client_returning(_HttpError(status))
        with pytest.raises(AppError) as exc:
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert fragment in exc.value.message

    @pytest.mark.asyncio
    async def test_missing_key_is_called_out_explicitly(self, parsed, monkeypatch) -> None:
        monkeypatch.setattr(settings, "ai_api_key", "")
        with pytest.raises(AppError) as exc:
            await OpenAIProvider().generate_solutions(parsed)
        assert "not configured" in exc.value.message


def _student() -> StudentInfo:
    return StudentInfo(
        name="Haziq",
        roll_number="242242",
        university=University.NUST,
        class_section="BSCS-3A",
        instructor_name="Dr. X",
        course="PF",
    )


def _all_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    return "\n".join(
        "".join(t.text or "" for t in p.xpath(".//w:t")) for p in document.element.xpath(".//w:p")
    )


class TestReportContent:
    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_report_contains_objectives_explanation_conclusion_and_class_section(self, university) -> None:
        student = _student().model_copy(update={"university": university})
        lab = GeneratedLab.model_validate(_LAB)
        text = _all_text(DocxGenerator().generate(student, lab))

        assert "Objectives" in text and "Learn loops" in text
        assert "Explanation:" in text and "Prints one." in text
        assert "Conclusion" in text and "Loops are useful." in text
        assert "BSCS-3A" in text  # the form field used to be silently dropped
        assert "{{" not in text  # no unfilled placeholders left behind

    def test_empty_objectives_and_conclusion_add_no_empty_headings(self) -> None:
        lab = GeneratedLab(
            lab_title="L",
            objectives=[],
            tasks=[
                GeneratedTaskSolution(
                    id="t", description="d", language="", filename="", code="", explanation=""
                )
            ],
            conclusion="",
        )
        text = _all_text(DocxGenerator().generate(_student(), lab))
        assert "Objectives" not in text and "Conclusion" not in text and "Explanation:" not in text


class TestConfig:
    def test_cors_origins_accepts_a_comma_separated_list(self) -> None:
        cfg = Settings(frontend_origin="https://a.example/, http://localhost:5173 ,", _env_file=None)
        assert cfg.cors_origins == ["https://a.example", "http://localhost:5173"]


class TestStaticServing:
    def test_built_frontend_is_served_and_api_still_wins(self, tmp_path, monkeypatch) -> None:
        import importlib

        (tmp_path / "index.html").write_text("<html>SPA</html>")
        monkeypatch.setattr(settings, "frontend_dist_dir", str(tmp_path))
        import app.main as main_module

        reloaded = importlib.reload(main_module)
        try:
            from fastapi.testclient import TestClient

            client = TestClient(reloaded.app)
            assert "SPA" in client.get("/").text
            assert client.get("/api/v1/health").status_code == 200
        finally:
            monkeypatch.setattr(settings, "frontend_dist_dir", "")
            importlib.reload(main_module)


class TestDocxTablesAreParsed:
    def _docx(self, build) -> bytes:
        document = Document()
        build(document)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def test_table_rows_appear_in_document_order(self) -> None:
        from app.parsing.docx_parser import parse_docx

        def build(d):
            d.add_paragraph("Before.")
            table = d.add_table(rows=2, cols=3)
            for i, row in enumerate([["Screen area", "Layout to use", "Lab"], ["Header", "RelativeLayout", "03"]]):
                for j, value in enumerate(row):
                    table.cell(i, j).text = value
            d.add_paragraph("After.")

        text = parse_docx(self._docx(build))
        assert text.splitlines() == ["Before.", "Screen area | Layout to use | Lab", "Header | RelativeLayout | 03", "After."]

    def test_merged_cells_are_not_repeated(self) -> None:
        from app.parsing.docx_parser import parse_docx

        def build(d):
            table = d.add_table(rows=1, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"

        assert parse_docx(self._docx(build)).splitlines() == ["Merged heading"]

    def test_document_with_only_a_table_is_not_empty(self) -> None:
        from app.parsing.docx_parser import parse_docx

        def build(d):
            d.add_table(rows=1, cols=1).cell(0, 0).text = "Only content"

        assert parse_docx(self._docx(build)) == "Only content"


def _response(content, finish_reason="stop", completion=100, reasoning=10):
    """A fake chat completion with finish_reason and usage like Groq returns."""
    message = MagicMock()
    message.content = content
    choice = MagicMock()
    choice.message = message
    choice.finish_reason = finish_reason
    response = MagicMock()
    response.choices = [choice]
    response.usage.completion_tokens = completion
    response.usage.completion_tokens_details.reasoning_tokens = reasoning
    return response


def _scripted_client(*results):
    """Each call returns/raises the next scripted result; records kwargs."""
    calls = []

    async def create(**kwargs):
        calls.append(kwargs)
        item = results[min(len(calls), len(results)) - 1]
        if isinstance(item, Exception):
            raise item
        return item

    client = MagicMock()
    client.chat.completions.create = create
    client.calls = calls
    return client


class TestReasoningModelBudget:
    """Regression: Groq openai/gpt-oss-120b returned empty content because the
    request carried no output budget and spent it all on hidden reasoning."""

    @pytest.fixture
    def parsed(self) -> ParsedLab:
        return ParsedLab(title="t", raw_text="Task 1: print", tasks=[])

    @pytest.mark.asyncio
    async def test_request_carries_an_explicit_budget_and_low_effort_for_gpt_oss(self, parsed, monkeypatch) -> None:
        monkeypatch.setattr(settings, "ai_max_output_tokens", 9000)
        monkeypatch.setattr(settings, "ai_reasoning_effort", "")
        client = _scripted_client(_response(json.dumps(_LAB)))
        await OpenAIProvider(client=client, model="openai/gpt-oss-120b").generate_solutions(parsed)
        sent = client.calls[0]
        assert sent["max_completion_tokens"] == 9000
        assert sent["extra_body"] == {"reasoning_effort": "low"}
        assert sent["response_format"] == {"type": "json_object"}

    @pytest.mark.asyncio
    async def test_reasoning_effort_is_not_sent_to_other_models_unless_configured(self, parsed, monkeypatch) -> None:
        monkeypatch.setattr(settings, "ai_reasoning_effort", "")
        client = _scripted_client(_response(json.dumps(_LAB)))
        await OpenAIProvider(client=client, model="gpt-4o-mini").generate_solutions(parsed)
        assert "extra_body" not in client.calls[0]

    @pytest.mark.asyncio
    async def test_configured_effort_overrides_the_automatic_one(self, parsed, monkeypatch) -> None:
        monkeypatch.setattr(settings, "ai_reasoning_effort", "medium")
        client = _scripted_client(_response(json.dumps(_LAB)))
        await OpenAIProvider(client=client, model="openai/gpt-oss-20b").generate_solutions(parsed)
        assert client.calls[0]["extra_body"] == {"reasoning_effort": "medium"}

    @pytest.mark.asyncio
    async def test_empty_answer_cut_off_by_length_gives_an_actionable_message(self, parsed) -> None:
        client = _scripted_client(_response("", finish_reason="length"))
        with pytest.raises(AppError) as exc:
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert exc.value.code == "AI_GENERATION_FAILED"
        assert "ran out of output space" in exc.value.message

    @pytest.mark.asyncio
    async def test_empty_answer_without_length_is_reported_as_empty(self, parsed) -> None:
        client = _scripted_client(_response("  ", finish_reason="stop"))
        with pytest.raises(AppError) as exc:
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert "empty response" in exc.value.message

    @pytest.mark.asyncio
    async def test_truncated_json_is_reported_as_out_of_space_not_unexpected(self, parsed) -> None:
        cut_off = json.dumps(_LAB)[:60]  # JSON that stops mid-way
        client = _scripted_client(_response(cut_off, finish_reason="length"))
        with pytest.raises(AppError) as exc:
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert "ran out of output space" in exc.value.message

    @pytest.mark.asyncio
    async def test_request_falls_back_json_mode_then_minimal(self, parsed) -> None:
        client = _scripted_client(_HttpError(400), _HttpError(400), _response(json.dumps(_LAB)))
        lab = await OpenAIProvider(client=client, model="openai/gpt-oss-120b").generate_solutions(parsed)
        assert lab.lab_title == "Lab 1"
        first, second, third = client.calls
        assert "response_format" in first
        assert "response_format" not in second and "max_completion_tokens" in second
        assert "max_tokens" in third and "extra_body" not in third and "max_completion_tokens" not in third

    @pytest.mark.asyncio
    async def test_non_400_errors_do_not_trigger_fallbacks(self, parsed) -> None:
        client = _scripted_client(_HttpError(429))
        with pytest.raises(AppError):
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert len(client.calls) == 1

    @pytest.mark.asyncio
    async def test_request_too_large_has_its_own_message(self, parsed) -> None:
        client = _scripted_client(_HttpError(413))
        with pytest.raises(AppError) as exc:
            await OpenAIProvider(client=client, model="m").generate_solutions(parsed)
        assert "per-minute token limit" in exc.value.message
