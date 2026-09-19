import io

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.schemas.lab import GeneratedLab, GeneratedTaskSolution, ParsedLab


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(72, 800, "Lab 04: Constraint Satisfaction Problems")
    c.drawString(72, 780, "Objective: Model the N-Queens problem as a CSP.")
    c.drawString(72, 760, "Task 1: Formulate variables, domains, and constraints.")
    c.save()
    return buffer.getvalue()


@pytest.fixture
def empty_pdf_bytes() -> bytes:
    """A structurally valid PDF with no text layer (simulates a scan)."""
    buffer = io.BytesIO()
    canvas.Canvas(buffer).save()
    return buffer.getvalue()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Lab 03: Prolog Basics", level=1)
    document.add_paragraph("Objective: Understand facts and rules in Prolog.")
    document.add_paragraph("Task 1: Write a Prolog program for family relationships.")
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def empty_docx_bytes() -> bytes:
    buffer = io.BytesIO()
    Document().save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_parsed_lab() -> ParsedLab:
    return ParsedLab(
        title="Lab 03: Prolog Basics",
        raw_text=(
            "Lab 03: Prolog Basics\n"
            "Objective: Understand facts and rules in Prolog.\n"
            "Task 1: Write a Prolog program for family relationships."
        ),
        tasks=[],
    )


@pytest.fixture
def sample_generated_lab() -> GeneratedLab:
    return GeneratedLab(
        lab_title="Lab 03: Prolog Basics",
        objectives=["Understand facts and rules in Prolog."],
        tasks=[
            GeneratedTaskSolution(
                id="task-1",
                description="Write a Prolog program for family relationships.",
                language="prolog",
                filename="family.pl",
                code="parent(tom, bob).\nparent(bob, ann).",
                explanation="Defines parent facts to represent family relationships.",
            )
        ],
        conclusion="This lab covered basic Prolog facts and rules.",
    )
