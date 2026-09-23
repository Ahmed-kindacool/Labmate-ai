import base64
import io

import pytest
from docx import Document

from app.core.errors import AppError
from app.documents import DocxGenerator
from app.domain.university import University
from app.schemas.generate import TaskExecutionResult
from app.schemas.lab import ExecutionResult, ExecutionStatus, GeneratedLab, GeneratedTaskSolution
from app.schemas.student import StudentInfo


def _student(university: University = University.AIR) -> StudentInfo:
    return StudentInfo(
        name="Ahmed Ali Khan",
        roll_number="211-CS-101",
        university=university,
        class_section="BS-CS 5A",
        instructor_name="Dr. Farah",
        course="Artificial Intelligence",
    )


def _lab(num_tasks: int = 2) -> GeneratedLab:
    tasks = [
        GeneratedTaskSolution(
            id=f"task-{i}",
            description=f"Task {i} description.",
            language="python",
            filename=f"task{i}.py",
            code=f"print('task {i}')",
            explanation=f"Explains task {i}.",
        )
        for i in range(1, num_tasks + 1)
    ]
    return GeneratedLab(
        lab_title="Lab 05: Search Algorithms",
        objectives=["Understand BFS.", "Understand DFS."],
        tasks=tasks,
        conclusion="This lab covered graph search algorithms.",
    )


def _document_full_text(document: Document) -> str:
    chunks = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def _load(docx_bytes: bytes) -> Document:
    return Document(io.BytesIO(docx_bytes))


class TestBasicFieldSubstitution:
    def test_student_and_lab_fields_are_filled_in(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab())
        document = _load(result)
        text = _document_full_text(document)

        assert "Ahmed Ali Khan" in text
        assert "211-CS-101" in text
        assert "BS-CS 5A" in text
        assert "Dr. Farah" in text
        assert "Artificial Intelligence" in text
        assert "Lab 05: Search Algorithms" in text
        assert "This lab covered graph search algorithms." in text

    def test_no_placeholder_tokens_survive_in_the_output(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab())
        text = _document_full_text(_load(result))

        for token in (
            "{{STUDENT_NAME}}",
            "{{ROLL_NUMBER}}",
            "{{CLASS_SECTION}}",
            "{{INSTRUCTOR_NAME}}",
            "{{COURSE}}",
            "{{LAB_TITLE}}",
            "{{OBJECTIVES}}",
            "{{TASKS}}",
            "{{CODE}}",
            "{{OUTPUT_SCREENSHOT}}",
            "{{CONCLUSION}}",
        ):
            assert token not in text

    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_works_for_every_university(self, university: University) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(university), _lab())
        assert result[:2] == b"PK"  # a real .docx is a real zip archive
        text = _document_full_text(_load(result))
        assert "Ahmed Ali Khan" in text
        # Never insert the university name as report text (Phase 6's
        # "Critical requirement", still binding in Phase 7's output).
        for name in ("air university", "bahria university", "nust"):
            assert name not in text.lower()


class TestObjectivesAndTasks:
    def test_each_objective_becomes_its_own_bullet_paragraph(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab())
        document = _load(result)

        bullets = [p.text for p in document.paragraphs if p.text.startswith("\u2022")]
        assert any("Understand BFS." in b for b in bullets)
        assert any("Understand DFS." in b for b in bullets)

    def test_empty_objectives_get_a_fallback_note_not_a_crash(self) -> None:
        lab = _lab()
        lab.objectives = []
        generator = DocxGenerator()
        result = generator.generate(_student(), lab)
        text = _document_full_text(_load(result))
        assert "No objectives were extracted" in text

    def test_each_task_gets_a_numbered_description(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab(num_tasks=3))
        text = _document_full_text(_load(result))

        assert "Task 1:" in text
        assert "Task 2:" in text
        assert "Task 3:" in text
        assert "Task 1 description." in text
        assert "Task 3 description." in text

    def test_no_tasks_gets_a_fallback_note_not_a_crash(self) -> None:
        lab = _lab(num_tasks=0)
        generator = DocxGenerator()
        result = generator.generate(_student(), lab)
        text = _document_full_text(_load(result))
        assert "No tasks were extracted" in text


class TestCodeRendering:
    def test_each_tasks_code_appears_under_its_own_filename_heading(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab(num_tasks=2))
        text = _document_full_text(_load(result))

        assert "Task 1 \u2014 task1.py" in text
        assert "print('task 1')" in text
        assert "Task 2 \u2014 task2.py" in text
        assert "print('task 2')" in text

    def test_multiline_code_preserves_every_line(self) -> None:
        lab = _lab(num_tasks=1)
        lab.tasks[0].code = "def add(a, b):\n    return a + b\n\nprint(add(2, 3))"
        generator = DocxGenerator()
        result = generator.generate(_student(), lab)
        document = _load(result)

        # Find the paragraph holding the code and check every source line
        # is present in its rendered text (line breaks collapse to "" in
        # paragraph.text, but the text content itself must survive).
        code_paragraph_text = "".join(
            p.text for p in document.paragraphs if "def add" in p.text
        )
        for line in ("def add(a, b):", "return a + b", "print(add(2, 3))"):
            assert line in code_paragraph_text

    def test_task_with_no_code_shows_a_note_instead_of_blank_space(self) -> None:
        lab = _lab(num_tasks=1)
        lab.tasks[0].code = ""
        generator = DocxGenerator()
        result = generator.generate(_student(), lab)
        text = _document_full_text(_load(result))
        assert "no code required" in text.lower()


class TestOutputRendering:
    def test_real_screenshot_is_embedded_as_an_image(self) -> None:
        lab = _lab(num_tasks=1)
        task_id = lab.tasks[0].id
        # A minimal but real 1x1 PNG, so add_picture has real image bytes
        # to decode -- not just an arbitrary byte string.
        png_bytes = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
            "YAAAAAYAAjCB0C8AAAAASUVORK5CYII="
        )
        screenshots = {task_id: base64.b64encode(png_bytes).decode("ascii")}
        execution_results = [
            TaskExecutionResult(
                task_id=task_id,
                result=ExecutionResult(status=ExecutionStatus.SUCCESS, stdout="ok", stderr=""),
            )
        ]

        generator = DocxGenerator()
        result = generator.generate(_student(), lab, execution_results, screenshots)
        document = _load(result)

        # A real inline image was added to the document's related parts.
        image_parts = [rel for rel in document.part.rels.values() if "image" in rel.reltype]
        assert len(image_parts) >= 1

    def test_corrupt_screenshot_bytes_fall_back_to_raw_stdout_instead_of_crashing(self) -> None:
        # A real integration-boundary bug: if the "screenshot" bytes
        # aren't a real, decodable image (e.g. a stand-in fake used by an
        # upstream test double, or genuine corruption), the whole report
        # must not fail -- it should degrade the same way a missing
        # screenshot does.
        lab = _lab(num_tasks=1)
        task_id = lab.tasks[0].id
        screenshots = {task_id: base64.b64encode(b"\x89PNG-not-a-real-image").decode("ascii")}
        execution_results = [
            TaskExecutionResult(
                task_id=task_id,
                result=ExecutionResult(
                    status=ExecutionStatus.SUCCESS, stdout="real stdout survives", stderr=""
                ),
            )
        ]

        generator = DocxGenerator()
        result = generator.generate(_student(), lab, execution_results, screenshots)
        text = _document_full_text(_load(result))
        assert "real stdout survives" in text

    def test_missing_screenshot_falls_back_to_raw_stdout_text(self) -> None:
        lab = _lab(num_tasks=1)
        task_id = lab.tasks[0].id
        execution_results = [
            TaskExecutionResult(
                task_id=task_id,
                result=ExecutionResult(
                    status=ExecutionStatus.SUCCESS, stdout="hello from stdout", stderr=""
                ),
            )
        ]

        generator = DocxGenerator()
        # No screenshots dict passed at all.
        result = generator.generate(_student(), lab, execution_results)
        text = _document_full_text(_load(result))
        assert "hello from stdout" in text

    def test_failed_execution_shows_real_stderr_not_stdout(self) -> None:
        lab = _lab(num_tasks=1)
        task_id = lab.tasks[0].id
        execution_results = [
            TaskExecutionResult(
                task_id=task_id,
                result=ExecutionResult(
                    status=ExecutionStatus.FAILED,
                    stdout="",
                    stderr="Traceback: ZeroDivisionError",
                ),
            )
        ]

        generator = DocxGenerator()
        result = generator.generate(_student(), lab, execution_results)
        text = _document_full_text(_load(result))
        assert "Traceback: ZeroDivisionError" in text

    def test_unsupported_language_shows_an_honest_note(self) -> None:
        lab = _lab(num_tasks=1)
        task_id = lab.tasks[0].id
        execution_results = [
            TaskExecutionResult(
                task_id=task_id,
                result=ExecutionResult(status=ExecutionStatus.UNSUPPORTED, stdout="", stderr=""),
            )
        ]

        generator = DocxGenerator()
        result = generator.generate(_student(), lab, execution_results)
        text = _document_full_text(_load(result))
        assert "not supported" in text.lower()

    def test_task_never_executed_shows_an_honest_note_not_fabricated_output(self) -> None:
        # No execution_results entry at all for this task -- e.g. it had
        # no code to run in the first place.
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab(num_tasks=1), execution_results=[])
        text = _document_full_text(_load(result))
        assert "No code was executed for this task." in text

    def test_no_tasks_gets_a_fallback_note_not_a_crash(self) -> None:
        generator = DocxGenerator()
        result = generator.generate(_student(), _lab(num_tasks=0))
        text = _document_full_text(_load(result))
        assert "No tasks were executed" in text


class TestErrorHandling:
    def test_unknown_university_propagates_template_registrys_app_error(self, monkeypatch) -> None:
        # DocxGenerator must not swallow or re-wrap TemplateRegistry's own
        # AppError into a different/duplicate message.
        student = _student()
        monkeypatch.setattr(student, "university", "mit", raising=False)

        generator = DocxGenerator()
        with pytest.raises(AppError) as exc_info:
            # Call get_template directly via a student whose .university
            # bypassed Pydantic validation, to exercise this path without
            # relying on StudentInfo accepting an invalid enum value.
            generator._templates.get_template(student.university)
        assert exc_info.value.code == "REPORT_GENERATION_FAILED"

    def test_corrupt_template_file_raises_report_generation_failed(self, tmp_path) -> None:
        from app.templates_registry import TemplateRegistry

        uni_dir = tmp_path / "air"
        uni_dir.mkdir()
        (uni_dir / "template.docx").write_bytes(b"not a real docx file")
        (uni_dir / "logo.png").write_bytes(b"not a real png")

        generator = DocxGenerator(template_registry=TemplateRegistry(templates_root=tmp_path))
        with pytest.raises(AppError) as exc_info:
            generator.generate(_student(), _lab())
        assert exc_info.value.code == "REPORT_GENERATION_FAILED"
