"""Project / documentation assignments (e.g. an Android Studio lab):
source files as real text, steps, filled tables and screenshot frames."""

import io
import json

import pytest
from docx import Document

from app.ai.openai_provider import _parse_generated_lab
from app.ai.prompts import SYSTEM_PROMPT
from app.documents import DocxGenerator
from app.domain.university import University
from app.execution import is_runnable_language
from app.schemas.lab import GeneratedLab
from app.schemas.student import StudentInfo

XML = '<?xml version="1.0" encoding="utf-8"?>\n<RelativeLayout android:id="@+id/rootLayout">\n    <TextView/>\n</RelativeLayout>'
JAVA = "public class MainActivity {\n    void onCreate() {\n        // body\n    }\n}"

ANDROID_LAB = {
    "lab_title": "Quiz Assignment",
    "objectives": ["Build the screen"],
    "tasks": [
        {"id": "task-1", "description": "Layout", "language": "xml", "filename": "activity_main.xml", "code": XML, "explanation": "The layout."},
        {"id": "task-2", "description": "Activity", "language": "java", "filename": "MainActivity.java", "code": JAVA, "explanation": "The activity."},
    ],
    "conclusion": "Done.",
    "steps": [
        {"title": "Create project", "details": "New project QuizApp", "location": "File > New", "issue": "Sync fails: re-sync."},
        "Run the app",
    ],
    "tables": [
        {"title": "Lifecycle Table", "headers": ["Action", "Callbacks"], "rows": [["Opened", "onCreate -> onStart -> onResume"], ["Back", "onPause -> onStop -> onDestroy"]]}
    ],
    "screenshots": ["Running app", "Toast"],
    "code_as_text": True,
}


def _student(university=University.AIR) -> StudentInfo:
    return StudentInfo(
        name="Haziq", roll_number="242242", university=university,
        class_section="BSSE-V-B", instructor_name="Miss Iqra", course="MAD",
    )


def _doc(lab: GeneratedLab, university=University.AIR):
    return Document(io.BytesIO(DocxGenerator().generate(_student(university), lab)))


def _text(document) -> str:
    return "\n".join("".join(t.text or "" for t in p.xpath(".//w:t")) for p in document.element.xpath(".//w:p"))


class TestRunnableLanguages:
    @pytest.mark.parametrize("language", ["python", "Python3", "py", "cpp", "c++", "C"])
    def test_runnable(self, language: str) -> None:
        assert is_runnable_language(language)

    @pytest.mark.parametrize("language", ["xml", "java", "kotlin", "html", "sql", ""])
    def test_source_only(self, language: str) -> None:
        assert not is_runnable_language(language)


class TestAndroidStyleReport:
    @pytest.fixture
    def document(self):
        return _doc(GeneratedLab.model_validate(_parse_generated_lab(json.dumps(ANDROID_LAB)).model_dump()))

    def test_every_required_section_is_present(self, document) -> None:
        text = _text(document)
        for expected in (
            "Steps Performed", "1. Create project", "2.", "Common issue and fix: Sync fails",
            "Task 1: activity_main.xml", "Task 2: MainActivity.java",
            "Lifecycle Table", "onPause -> onStop -> onDestroy",
            "Screenshots", "Figure 1: Running app", "Figure 2: Toast", "Conclusion",
        ):
            assert expected in text, expected

    def test_code_is_real_text_not_an_image(self, document) -> None:
        assert len(document.inline_shapes) == 0
        text = _text(document)
        assert 'android:id="@+id/rootLayout"' in text and "class MainActivity" in text

    def test_code_indentation_survives(self, document) -> None:
        assert "    <TextView/>" in _text(document)

    def test_no_output_blocks_or_placeholders_for_source_only_tasks(self, document) -> None:
        text = _text(document)
        assert "Execution Output" not in text
        assert "No execution output" not in text
        assert "No code provided" not in text
        assert "{{" not in text

    def test_screenshot_frames_are_empty_placeholders(self, document) -> None:
        assert _text(document).count("[ Insert your screenshot here ]") == 2

    @pytest.mark.parametrize("university", [University.AIR, University.BAHRIA, University.NUST])
    def test_cover_has_a_date_on_every_template(self, university) -> None:
        lab = GeneratedLab.model_validate(ANDROID_LAB | {"steps": [], "tables": [], "screenshots": []})
        assert "Date:" in _text(_doc(lab, university))

    def test_document_stays_a_valid_docx_with_the_new_tables(self, document, tmp_path) -> None:
        path = tmp_path / "x.docx"
        document.save(path)
        reopened = Document(path)  # round-trips without XML errors
        assert len(reopened.tables) >= 4


class TestBackwardsCompatibility:
    def test_plain_python_lab_json_still_parses_with_empty_new_fields(self) -> None:
        lab = _parse_generated_lab(json.dumps({"lab_title": "L", "tasks": [{"description": "d", "language": "python", "code": "print(1)"}]}))
        assert lab.steps == [] and lab.tables == [] and lab.screenshots == [] and lab.code_as_text is False

    def test_garbage_in_new_fields_is_ignored_not_fatal(self) -> None:
        raw = {"lab_title": "L", "tasks": [], "steps": "nope", "tables": [5, {"rows": "x"}], "screenshots": {"a": 1}, "code_as_text": "maybe"}
        lab = _parse_generated_lab(json.dumps(raw))
        assert lab.steps == [] and lab.tables == [] and lab.screenshots == [] and lab.code_as_text is False

    def test_string_booleans_and_dict_captions_are_coerced(self) -> None:
        lab = _parse_generated_lab(json.dumps({"lab_title": "L", "tasks": [], "code_as_text": "true", "screenshots": [{"caption": "Logcat"}, ""]}))
        assert lab.code_as_text is True and lab.screenshots == ["Logcat"]

    def test_runnable_code_is_still_an_image_when_text_mode_is_off(self) -> None:
        lab = GeneratedLab.model_validate({
            "lab_title": "L", "objectives": [], "conclusion": "",
            "tasks": [{"id": "t", "description": "d", "language": "python", "filename": "a.py", "code": "print(1)", "explanation": ""}],
        })
        text = _text(_doc(lab))
        assert "Execution Output:" in text  # python tasks keep the output section


class TestPromptCoversProjectLabs:
    def test_prompt_asks_for_source_files_not_section_headings(self) -> None:
        assert "SOURCE FILE" in SYSTEM_PROMPT
        assert "Never turn the document's own headings" in SYSTEM_PROMPT
        assert '"steps"' in SYSTEM_PROMPT and '"tables"' in SYSTEM_PROMPT and '"screenshots"' in SYSTEM_PROMPT

    def test_prompt_no_longer_limits_languages_to_runnable_ones(self) -> None:
        assert "xml" in SYSTEM_PROMPT and "java" in SYSTEM_PROMPT

    def test_prompt_never_lets_steps_claim_personal_problems(self) -> None:
        assert "Never claim the student personally hit a problem" in SYSTEM_PROMPT
